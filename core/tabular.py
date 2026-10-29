"""Reading uploaded spreadsheets and writing Excel downloads.

Coordinators can upload either CSV or Excel (.xlsx). Both are turned into the same list of
rows, so every import works the same way whatever the file type.
"""
import csv
import io

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ACCEPTED = (".csv", ".xlsx")


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_rows(upload):
    """Return (column names, rows as dicts) from an uploaded .csv or .xlsx file.
    Column names are lower-cased with spaces turned into underscores, so "Full Name" and
    "full_name" both work."""
    name = upload.name.lower()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(upload.read()), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None) or []
        cols = [_cell(h).lower().replace(" ", "_") for h in header]
        rows = []
        for values in it:
            values = list(values) + [None] * (len(cols) - len(values))
            rows.append({c: _cell(v) for c, v in zip(cols, values) if c})
        wb.close()
    else:
        text = upload.read().decode("utf-8-sig", "replace")
        reader = csv.DictReader(io.StringIO(text))
        cols = [(f or "").strip().lower().replace(" ", "_") for f in (reader.fieldnames or [])]
        rows = [{c: (v or "").strip() for c, v in zip(cols, r.values()) if c} for r in reader]
    return [c for c in cols if c], rows


def xlsx_response(filename, title, headers, rows, widths=None):
    """An Excel file download with a bold, frozen header row."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="007A49")
    for c in ws[1]:
        c.font, c.fill = head, fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(60, max(10, longest + 2)))
    out = io.BytesIO()
    wb.save(out)
    resp = HttpResponse(out.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp
