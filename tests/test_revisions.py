"""Supervisor's second review: Excel and CSV imports for students and supervisors, Excel
downloads, and a coordinator for each department who runs that department's allocation
without touching anyone else's records."""
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from core import demo
from core.models import HND, ND, Allocation, Department, PastProject, Student, Supervisor, User

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx(name, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile(name, out.getvalue(), XLSX)


class Revisions(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", verbosity=0, no_eval=True)
        cls.cs = Department.objects.get(name="Computer Science")
        cls.ee = Department.objects.get(name="Electrical and Electronic Engineering")

    def setUp(self):
        self.client.login(username="coordinator@yabatech.demo", password=demo.PASSWORD)

    def test_student_import_from_excel(self):
        f = xlsx("students.xlsx", ["Matric No", "Full Name", "CGPA", "Interest Text"], [
            ["F/HD/24/3218001", "Excel Person", 3.4, "a web portal for booking hostel rooms online"],
            ["F/ND/24/3218002", "Excel ND Person", 2.9, "an app that lists the bus routes around Yaba"],
            ["F/HD/24/3218003", "Bad CGPA", 4.6, "a system for managing church members"],
            [None, None, None, None]])
        r = self.client.post(reverse("c_students_import"), {"file": f}, follow=True)
        self.assertContains(r, "Imported 2 student records")
        self.assertContains(r, "1 rows were skipped")
        self.assertEqual(float(Student.objects.get(pk="F/HD/24/3218001").cgpa), 3.4)
        self.assertEqual(Student.objects.get(pk="F/ND/24/3218002").level, ND)

    def test_old_excel_and_other_files_are_refused_kindly(self):
        r = self.client.post(reverse("c_students_import"),
                             {"file": SimpleUploadedFile("old.xls", b"x")}, follow=True)
        self.assertContains(r, "Excel Workbook (.xlsx)")
        r = self.client.post(reverse("c_students_import"),
                             {"file": SimpleUploadedFile("broken.xlsx", b"not a workbook", XLSX)}, follow=True)
        self.assertContains(r, "could not be read")

    def test_supervisor_import_adds_updates_never_deletes(self):
        before = self.cs.supervisors.count()
        existing = self.cs.supervisors.get(full_name="Dr. Halima Bello")
        f = xlsx("sups.xlsx", ["full_name", "specialisation", "hnd_capacity", "nd_capacity"], [
            ["Dr. New Lecturer", "Software engineering, testing and quality assurance", 4, 2],
            ["dr. halima bello", "Cybersecurity, digital forensics and secure software", 6, 3],
            ["No Places", "Mobile computing and wireless systems", 0, 0],
            ["Short Spec", "AI", 3, 3]])
        r = self.client.post(reverse("c_supervisors_import"), {"file": f}, follow=True)
        self.assertContains(r, "Added 1 and updated 1 supervisors")
        self.assertContains(r, "2 rows were skipped")
        self.assertEqual(self.cs.supervisors.count(), before + 1)
        existing.refresh_from_db()
        self.assertEqual((existing.capacity, existing.capacity_nd), (6, 3))
        self.assertTrue(existing.past_projects.exists())  # profile history kept
        csv_file = SimpleUploadedFile("s.csv", b"full_name,specialisation,hnd_capacity,nd_capacity\n"
                                                 b"Mrs. Csv Person,Data analysis and statistics for business,3,\n")
        self.client.post(reverse("c_supervisors_import"), {"file": csv_file})
        self.assertEqual(self.cs.supervisors.get(full_name="Mrs. Csv Person").capacity_nd, 3)  # default

        # the administrator gives an imported supervisor a sign-in: linked to the same record
        admin = User.objects.create_superuser(username="boss@example.com", email="boss@example.com",
                                              password="boss-secret-99")
        self.client.force_login(admin)
        r = self.client.post("/admin/core/user/add/", {
            "email": "new.lect@example.com", "first_name": "Dr. New", "last_name": "Lecturer",
            "role": "supervisor", "department": self.cs.pk, "password1": "lect-secret-77",
            "password2": "lect-secret-77", "usable_password": "true"})
        self.assertEqual(r.status_code, 302)
        sup = self.cs.supervisors.get(full_name="Dr. New Lecturer")
        self.assertEqual(sup.user.email, "new.lect@example.com")
        self.assertEqual(self.cs.supervisors.filter(full_name="Dr. New Lecturer").count(), 1)
        # and a brand-new supervisor added in the admin panel gets a record in their department
        self.client.post("/admin/core/user/add/", {
            "email": "brand.new@example.com", "first_name": "Mrs. Brand", "last_name": "New",
            "role": "supervisor", "department": self.cs.pk, "password1": "lect-secret-77",
            "password2": "lect-secret-77", "usable_password": "true"})
        self.assertTrue(self.cs.supervisors.filter(full_name="Mrs. Brand New", user__email="brand.new@example.com").exists())
        r = self.client.post("/admin/core/user/add/", {
            "email": "nodept@example.com", "role": "supervisor", "department": "",
            "password1": "lect-secret-77", "password2": "lect-secret-77", "usable_password": "true"})
        self.assertContains(r, "Choose the supervisor")

    def test_past_projects_from_excel(self):
        f = xlsx("p.xlsx", ["supervisor_name", "title", "keywords", "abstract", "session"], [
            ["Dr. Halima Bello", "Ransomware detection using file entropy analysis", "security", "", "2024/2025"]])
        r = self.client.post(reverse("c_projects"), {"file": f}, follow=True)
        self.assertContains(r, "Loaded 1 past project records")
        self.assertTrue(PastProject.objects.filter(title__startswith="Ransomware").exists())

    def test_excel_downloads(self):
        self.client.get(reverse("c_overview") + "?level=HND")
        self.client.post(reverse("c_run"), {"threshold": "0.30"})
        for name in ["c_report_xlsx", "c_students_xlsx", "c_supervisors_xlsx"]:
            r = self.client.get(reverse(name))
            self.assertEqual(r["Content-Type"], XLSX, name)
            ws = load_workbook(io.BytesIO(r.content)).active
            self.assertGreater(ws.max_row, 5, name)
        ws = load_workbook(io.BytesIO(self.client.get(reverse("c_report_xlsx")).content)).active
        self.assertEqual(ws.max_row, self.cs.students.filter(level=HND).count() + 1)
        for t in ["students_template", "supervisors_template", "past_projects_template"]:
            self.assertEqual(self.client.get(reverse("xlsx_template", args=[t]))["Content-Type"], XLSX)
            self.assertEqual(self.client.get(reverse("csv_template", args=[t])).status_code, 200)

    def test_each_department_runs_its_own_allocation(self):
        # Electrical has already allocated and published; Computer Science then runs twice
        self.client.logout()
        self.client.login(username="coordinator.ee@yabatech.demo", password=demo.PASSWORD)
        self.client.get(reverse("c_overview") + "?level=HND")
        self.client.post(reverse("c_run"), {"threshold": "0.0"})
        self.client.post(reverse("c_report"), {"publish": 1})
        ee_before = list(Allocation.objects.filter(student__department=self.ee)
                         .values_list("student_id", "supervisor_id"))
        self.assertTrue(self.ee.cycle(HND).published)
        n_students = Student.objects.count()

        self.client.logout()
        self.client.login(username="coordinator@yabatech.demo", password=demo.PASSWORD)
        self.client.get(reverse("c_overview") + "?level=HND")
        self.client.post(reverse("c_run"), {"threshold": "0.30"})
        self.client.post(reverse("c_run"), {"threshold": "0.25"})
        self.assertEqual(Student.objects.count(), n_students)  # no student record deleted
        self.assertEqual(list(Allocation.objects.filter(student__department=self.ee)
                              .values_list("student_id", "supervisor_id")), ee_before)
        self.assertTrue(Department.objects.get(pk=self.ee.pk).cycle(HND).published)

    def test_coordinators_are_created_only_by_the_administrator(self):
        self.client.logout()
        stats = Department.objects.get(name="Statistics")
        self.client.post(reverse("register"), {"role": "coordinator", "department": stats.pk,
                                               "full_name": "Mrs. Stats Coordinator",
                                               "email": "stats@example.com", "password": "a-long-pass-1"})
        self.assertFalse(User.objects.filter(email="stats@example.com").exists())
        admin = User.objects.create_superuser(username="boss@example.com", email="boss@example.com",
                                              password="boss-secret-99")
        self.client.force_login(admin)
        r = self.client.post("/admin/core/user/add/", {
            "email": "stats@example.com", "first_name": "Mrs. Stats", "last_name": "Coordinator",
            "role": "coordinator", "department": stats.pk, "password1": "stats-secret-77",
            "password2": "stats-secret-77", "usable_password": "true"})
        self.assertEqual(r.status_code, 302)
        self.client.logout()
        self.assertTrue(self.client.login(username="stats@example.com", password="stats-secret-77"))
        r = self.client.get(reverse("c_overview"))
        self.assertContains(r, "Department of Statistics")
        self.assertContains(r, "No HND allocation yet")
        self.assertNotContains(self.client.get(reverse("c_students")), "F/HD/24/32100")  # sees only its own
